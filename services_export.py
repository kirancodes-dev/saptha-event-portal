"""
services_export.py — one export of an event's registrations, in CSV and
Excel, and a one-click event report (UPG-03).

Used by the SPOC, coordinator and admin exports. Callers check
can(user, 'export_data', event) first; nothing here checks permissions.
"""
import csv
import datetime
import io
import re
from collections import Counter, defaultdict

# One row per registration; a team's lead, then its other members
COLUMNS = ['Ticket ID', 'Lead Name', 'Lead USN', 'Department', 'Year', 'Lead Email', 'Lead Phone', 'Team', 'Members',
           'Attendance', 'Check-in Time', 'Payment Status', 'Amount (Rs)', 'Score', 'Rank',
           'Certificate ID', 'Registered At', 'Participant Type', 'Designation', 'Institution']

def _ptype(record):
    """Student, faculty or external (UPG-12); older records are students."""
    from services_participants import type_of
    return type_of(record)


USN = re.compile(r'^\d?[A-Z]{2}(\d{2})([A-Z]{2,3})\d+$')


def _answers(reg):
    answers = reg.get('form_answers')
    return answers if isinstance(answers, dict) else {}


def department_and_year(reg, profile=None):
    """Department: the profile's, else the form's, else the USN's branch code.
    Year: the form's year or semester answer, else the batch in the USN
    (1SN22CS007 → 2022)."""
    answers = _answers(reg)
    usn = str(reg.get('lead_usn') or answers.get('usn') or '').strip().upper()
    m = USN.match(usn)
    department = ((profile or {}).get('department') or answers.get('department') or answers.get('branch')
                  or (m.group(2) if m else ''))
    year = answers.get('year') or answers.get('year_of_study') or answers.get('semester') or (
        f"20{m.group(1)}" if m else '')
    return str(department or ''), str(year or '')


def _average(reg):
    totals = [s.get('total') for s in (reg.get('scores') or {}).values() if isinstance(s, dict)]
    totals = [float(t) for t in totals if isinstance(t, (int, float)) or str(t).replace('.', '', 1).isdigit()]
    return round(sum(totals) / len(totals), 2) if totals else None


def _ranks(regs):
    """Published ranks where there are any; otherwise by average score."""
    if any(r.get('final_rank') for _, r in regs):
        return {rid: r.get('final_rank') for rid, r in regs if r.get('final_rank')}
    scored = sorted(((rid, _average(r)) for rid, r in regs if _average(r) is not None), key=lambda x: -x[1])
    return {rid: i for i, (rid, _) in enumerate(scored, start=1)}


def _registrations(db, event_id):
    return [(d.id, d.to_dict() or {}) for d in db.collection('registrations').where('event_id', '==', event_id).stream()]


def _profiles(db, regs):
    out = {}
    for _, reg in regs:
        email = (reg.get('lead_email') or '').lower()
        if email and email not in out:
            doc = db.collection('users').document(email).get()
            out[email] = (doc.to_dict() or {}) if doc.exists else {}
    return out


def registration_rows(db, event_id):
    """One row per registration, in COLUMNS order."""
    regs = _registrations(db, event_id)
    profiles = _profiles(db, regs)
    ranks = _ranks(regs)
    rows = []
    for rid, r in sorted(regs, key=lambda x: str(x[1].get('registered_at') or '')):
        department, year = department_and_year(r, profiles.get((r.get('lead_email') or '').lower()))
        others = [m for m in (r.get('members') or []) if (m.get('email') or '') != r.get('lead_email')]
        score = r.get('final_score') if r.get('final_score') not in (None, '') else _average(r)
        rows.append([
            r.get('reg_id') or rid, r.get('lead_name', ''), r.get('lead_usn', ''), department, year,
            r.get('lead_email', ''), r.get('lead_phone') or r.get('phone', ''),
            r.get('team_name', ''), '; '.join(f"{m.get('name', '')} ({m.get('usn', '')})" for m in others),
            r.get('attendance', 'Pending'), r.get('checkin_time', ''),
            r.get('payment_status', ''), r.get('amount_paid', 0) or 0,
            '' if score is None else score, ranks.get(rid, ''),
            r.get('certificate_id', ''), r.get('registered_at', ''),
            _ptype(r), r.get('designation', ''), r.get('institution', ''),   # UPG-12
        ])
    return rows


def to_csv(rows, columns=COLUMNS):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(columns)
    for row in rows:
        writer.writerow([_cell(v) for v in row])
    return out.getvalue()


def _cell(value):
    # A spreadsheet runs a cell that starts with = + - @ as a formula, and
    # registrants type names: quote those, but leave numbers like +91 98… alone
    if (isinstance(value, str) and value[:1] in ('=', '+', '-', '@', '\t', '\r')
            and not re.fullmatch(r'[+-]?[\d\s().-]+', value)):
        return "'" + value
    return value


def to_xlsx(sheets):
    """sheets: [(title, columns, rows)] → .xlsx bytes."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, columns, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='0D2D62')
        for row in rows:
            ws.append([_cell(v) for v in row])
        for i, name in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(str(name)) + 4))
        ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def registrations_xlsx(db, event_id):
    return to_xlsx([('Registrations', COLUMNS, registration_rows(db, event_id))])


def all_registration_rows(db):
    """Every event's registrations (the admin export), event title first."""
    rows = []
    for e in db.collection('events').stream():
        title = (e.to_dict() or {}).get('title', '')
        rows += [[title] + row for row in registration_rows(db, e.id)]
    return rows


# ── Event report ───────────────────────────────────────────────────────────

def event_report(db, event_id):
    """Registrations against attendance by department and year, feedback
    averages and winners. No AI."""
    ev_doc = db.collection('events').document(event_id).get()
    event = (ev_doc.to_dict() or {}) if ev_doc.exists else {}
    regs = _registrations(db, event_id)
    profiles = _profiles(db, regs)
    ranks = _ranks(regs)

    by_dept, by_year, by_type = defaultdict(Counter), defaultdict(Counter), defaultdict(Counter)
    ratings = []
    for rid, r in regs:
        department, year = department_and_year(r, profiles.get((r.get('lead_email') or '').lower()))
        present = r.get('attendance') == 'Present'
        for table, key in ((by_dept, department or 'Not given'), (by_year, year or 'Not given'),
                           (by_type, _ptype(r))):
            table[key]['registered'] += 1
            table[key]['present'] += 1 if present else 0
        rating = (r.get('feedback') or {}).get('rating') if isinstance(r.get('feedback'), dict) else None
        if isinstance(rating, (int, float)) or str(rating or '').isdigit():
            ratings.append(float(rating))

    winners = sorted(((ranks[rid], r) for rid, r in regs if rid in ranks and ranks[rid] <= 3), key=lambda x: x[0])
    return {
        'event': {'id': event_id, 'title': event.get('title', 'Event'), 'date': str(event.get('date', '')),
                  'venue': event.get('venue', '')},
        'registered': len(regs),
        'present': sum(1 for _, r in regs if r.get('attendance') == 'Present'),
        'by_department': {k: dict(v) for k, v in sorted(by_dept.items())},
        'by_year': {k: dict(v) for k, v in sorted(by_year.items())},
        'by_participant_type': {k: dict(v) for k, v in sorted(by_type.items())},   # UPG-12
        'feedback': {'responses': len(ratings),
                     'average_rating': round(sum(ratings) / len(ratings), 2) if ratings else None},
        'winners': [{'rank': rank, 'name': r.get('lead_name', ''), 'team': r.get('team_name', ''),
                     'score': r.get('final_score') or _average(r)} for rank, r in winners],
        'generated_at': datetime.datetime.now().strftime('%Y-%m-%d %H:%M'),
    }


def event_report_xlsx(db, event_id):
    report = event_report(db, event_id)
    ev = report['event']
    present_pct = round(100 * report['present'] / report['registered'], 1) if report['registered'] else 0
    summary = [
        ['Event', ev['title']], ['Date', ev['date']], ['Venue', ev['venue']],
        ['Registered', report['registered']], ['Present', report['present']],
        ['Attendance %', present_pct],
        ['Feedback responses', report['feedback']['responses']],
        ['Average feedback rating', '' if report['feedback']['average_rating'] is None
         else report['feedback']['average_rating']],
        ['Generated', report['generated_at']],
    ]
    breakdown = lambda table: [[k, v.get('registered', 0), v.get('present', 0)] for k, v in table.items()]  # noqa: E731
    return to_xlsx([
        ('Summary', ['Item', 'Value'], summary),
        ('By department', ['Department', 'Registered', 'Present'], breakdown(report['by_department'])),
        ('By year', ['Year', 'Registered', 'Present'], breakdown(report['by_year'])),
        ('By participant type', ['Participant type', 'Registered', 'Present'],
         breakdown(report['by_participant_type'])),
        ('Winners', ['Rank', 'Name', 'Team', 'Score'],
         [[w['rank'], w['name'], w['team'], '' if w['score'] is None else w['score']] for w in report['winners']]),
        ('Registrations', COLUMNS, registration_rows(db, event_id)),
    ])


# ── Department Activity & Hours Export (UPG-11) ───────────────────────────

DEPARTMENT_ACTIVITY_COLUMNS = [
    'Student Name', 'USN', 'Email', 'Department', 'Year',
    'Events Attended', 'Total Activity Points', 'Total Activity Hours', 'Participant Type',
]


def department_activity_data(db, department=None, participant_type=None):
    """Return list of dicts with activity totals for every student in the department."""
    students = []
    for u in db.collection('users').stream():
        udata = u.to_dict() or {}
        role = str(udata.get('role', ''))
        if role in ('Student', 'Participant'):
            stud_dept = str(udata.get('department') or '').strip()
            if department and stud_dept.lower() != str(department).strip().lower():
                continue
            # Faculty FDP hours too; ?participant_type= narrows to one type (UPG-12)
            if participant_type and _ptype(udata) != str(participant_type).strip().lower():
                continue
            students.append((u.id, udata))

    # Pre-fetch events
    events_map = {str(e.id): (e.to_dict() or {}) for e in db.collection('events').stream()}

    # Group registrations by email and usn
    regs_by_student = defaultdict(list)
    for r in db.collection('registrations').stream():
        rdata = r.to_dict() or {}
        email = str(rdata.get('lead_email') or '').strip().lower()
        usn = str(rdata.get('lead_usn') or '').strip().upper()
        if email:
            regs_by_student[email].append(rdata)
        if usn:
            regs_by_student[usn].append(rdata)
        for m in (rdata.get('members') or []):
            m_email = str(m.get('email') or '').strip().lower()
            m_usn = str(m.get('usn') or '').strip().upper()
            if m_email and m_email != email:
                regs_by_student[m_email].append(rdata)
            if m_usn and m_usn != usn:
                regs_by_student[m_usn].append(rdata)

    results = []
    for u_id, sdata in sorted(students, key=lambda x: (x[1].get('department', ''), x[1].get('name', ''))):
        s_email = str(sdata.get('email') or u_id or '').strip()
        s_usn = str(sdata.get('usn') or '').strip().upper()
        # Collect distinct registrations
        seen_reg_ids = set()
        matched_regs = []
        for r in regs_by_student.get(s_email.lower(), []) + (regs_by_student.get(s_usn, []) if s_usn else []):
            rid = str(r.get('id') or r.get('reg_id') or id(r))
            if rid not in seen_reg_ids:
                seen_reg_ids.add(rid)
                matched_regs.append(r)

        attended_count = 0
        total_points = 0.0
        total_hours = 0.0
        from services_ledger import credit_for
        for r in matched_regs:   # each registration once (deduplicated above)
            if str(r.get('attendance') or '').strip().lower() == 'present':
                attended_count += 1
                points, hours = credit_for(r, events_map.get(str(r.get('event_id')), {}))
                total_points += points
                total_hours += hours

        results.append({
            'name': sdata.get('name', ''),
            'usn': sdata.get('usn', ''),
            'email': s_email,
            'department': sdata.get('department', ''),
            'year': sdata.get('year', ''),
            'events_attended': attended_count,
            'total_points': total_points,
            'total_hours': total_hours,
            'participant_type': _ptype(sdata),
        })
    return results


def department_activity_rows(db, department=None, participant_type=None):
    """List every participant in the department with totals as row tuples."""
    data = department_activity_data(db, department, participant_type)
    return [
        [
            d['name'], d['usn'], d['email'], d['department'], d['year'],
            d['events_attended'], d['total_points'], d['total_hours'], d['participant_type'],
        ]
        for d in data
    ]


def department_activity_csv(db, department=None, participant_type=None):
    return to_csv(department_activity_rows(db, department, participant_type), columns=DEPARTMENT_ACTIVITY_COLUMNS)


def department_activity_xlsx(db, department=None, participant_type=None):
    return to_xlsx([('Department Activity', DEPARTMENT_ACTIVITY_COLUMNS,
                     department_activity_rows(db, department, participant_type))])


# ── Placement drive rounds (UPG-57) ────────────────────────────────────────

PLACEMENT_ROUND_COLUMNS = ['Ticket ID', 'Name', 'USN', 'Email', 'Phone', 'Department', 'Year',
                           'CGPA', 'Active Backlogs', 'Round', 'Result']


def placement_round_rows(db, event_id, round_id, status=None):
    """Everyone with a result in one round of a placement drive (optionally one result)."""
    from services_placement import academic_record, settings
    ev_doc = db.collection('events').document(event_id).get()
    event = (ev_doc.to_dict() or {}) if ev_doc.exists else {}
    round_name = next((r['name'] for r in settings(event).get('rounds') or [] if r['id'] == round_id), round_id)
    regs = _registrations(db, event_id)
    profiles = _profiles(db, regs)
    rows = []
    for rid, r in sorted(regs, key=lambda x: str(x[1].get('lead_name') or '')):
        result = (r.get('placement_rounds') or {}).get(round_id)
        if not result or (status and result != status):
            continue
        rec = academic_record(profiles.get((r.get('lead_email') or '').lower()) or {})
        rows.append([r.get('reg_id') or rid, r.get('lead_name', ''), r.get('lead_usn') or rec['usn'], r.get('lead_email', ''),
                     r.get('lead_phone', ''), rec['department'], rec['year'],
                     '' if rec['cgpa'] is None else rec['cgpa'], '' if rec['backlogs'] is None else rec['backlogs'],
                     round_name, result])
    return rows
